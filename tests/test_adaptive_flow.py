"""Exercise the real Streamlit submission path; normal tests never call an LLM."""

import json
import os
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock

import pytest
import streamlit as st
from docx import Document
from streamlit.testing.v1 import AppTest

from interview_coach import provider, recording, speech, transcription
from interview_coach.config import ModelCandidate, Settings
from interview_coach.provider import OpenAIQuestionProvider
from interview_coach.schemas import AnswerAnalysis, InterviewQuestion, InterviewQuestionSet
from interview_coach.transcription import transcribe_audio as real_transcribe_audio

APP = Path(__file__).resolve().parents[1] / "app.py"


def analysis(action="probe"):
    return AnswerAnalysis(
        action=action,
        answer_quality="partial",
        evidence_found="Mentions improvement.",
        missing_evidence="Personal actions and outcomes.",
        reasoning="Need specific evidence.",
        follow_up_question="What did you personally change?" if action != "advance" else None,
    )


@pytest.fixture
def flow(tmp_path, monkeypatch):
    questions = InterviewQuestionSet(
        summary="Assess ownership and impact.",
        questions=[
            InterviewQuestion(
                question=f"Describe an onboarding improvement you led. Example {i + 1}.",
                competency="Problem solving",
                difficulty="intermediate",
                source_focus="both",
                rationale="Assess ownership.",
                expected_evidence="Personal actions and results.",
            )
            for i in range(5)
        ],
    )
    model = Mock()
    model.generate.return_value = questions
    model.analyze_answer.return_value = analysis()
    monkeypatch.setattr(provider, "OpenAIQuestionProvider", lambda settings: model)
    monkeypatch.setattr(Settings, "from_env", lambda: Settings(candidates=()))
    transcribe = Mock(return_value="I improved the process.")
    monkeypatch.setattr(transcription, "transcribe_audio", transcribe)
    spoken = []
    monkeypatch.setattr(
        speech, "render_question_speech", lambda questions: spoken.append(questions)
    )
    real_save = recording.save_answer_recording

    def save(**kwargs):
        kwargs["root"] = tmp_path
        return real_save(**kwargs)

    monkeypatch.setattr(recording, "save_answer_recording", save)
    document = Document()
    document.add_paragraph("Product manager working on onboarding and payments. " * 4)
    resume = BytesIO()
    document.save(resume)
    upload = SimpleNamespace(name="resume.docx", getvalue=lambda: resume.getvalue())
    monkeypatch.setattr(st, "file_uploader", lambda *a, **kw: upload)
    audios = {}
    audio_keys = []

    def audio_input(*args, key, **kwargs):
        audio_keys.append(key)
        return BytesIO(audios[key]) if key in audios else None

    monkeypatch.setattr(st, "audio_input", audio_input)
    at = AppTest.from_file(str(APP), default_timeout=10).run()
    at.text_area[0].set_value("Hiring a product manager to improve onboarding and payments. " * 4)
    at.button[0].click().run()
    assert not at.exception

    def submit(payload=b"test audio", timeout=10):
        key = audio_keys[-1]
        audios[key] = payload
        at.run(timeout=timeout)
        at.button(key="submit_" + key.removeprefix("answer_audio_")).click().run(timeout=timeout)
        assert not at.exception
        return key

    def saved():
        return json.loads(next(tmp_path.glob("*/metadata.json")).read_text())["answers"]

    return SimpleNamespace(
        app=at,
        model=model,
        transcribe=transcribe,
        submit=submit,
        saved=saved,
        audios=audios,
        keys=audio_keys,
        spoken=spoken,
        root=tmp_path,
    )


@pytest.mark.parametrize("action", ["probe", "clarify", "challenge"])
def test_submission_displays_followup_and_clears_recording(flow, action):
    flow.model.analyze_answer.return_value = analysis(action)
    old_key = flow.submit()
    assert flow.app.session_state.current_question_index == 0
    assert flow.app.session_state.follow_up_count == 1
    assert "Follow-up 1 of 2" in [item.value for item in flow.app.subheader]
    assert any("What did you personally change?" in item.value for item in flow.app.markdown)
    assert flow.spoken[-1] == ["What did you personally change?"]
    assert flow.keys[-1] != old_key
    assert len(flow.app.button) == 1  # No previous audio to accidentally resubmit.
    assert flow.saved()[0]["analysis"]["action"] == action


def test_followup_limit_history_and_recordings(flow):
    flow.submit(b"original")
    flow.submit(b"followup one")
    flow.submit(b"followup two")  # Fake model keeps probing; service must enforce the limit.
    assert flow.app.session_state.current_question_index == 1
    assert flow.app.session_state.follow_up_count == 0
    rows = flow.saved()
    assert [r["follow_up_number"] for r in rows] == [0, 1, 2]
    assert rows[-1]["analysis"]["action"] == "advance"
    assert rows[1]["question"] == "What did you personally change?"
    assert len({r["audio_file"] for r in rows}) == 3
    directory = next(flow.root.glob("interview_*"))
    assert [(directory / r["audio_file"]).read_bytes() for r in rows] == [
        b"original",
        b"followup one",
        b"followup two",
    ]
    request = flow.model.analyze_answer.call_args.kwargs
    assert "Question:" in request["previous_context"]
    assert "Answer: I improved the process." in request["previous_context"]
    assert request["resume_text"] and request["job_description"]


def test_explicit_advance_moves_to_next_question(flow):
    flow.model.analyze_answer.return_value = analysis("advance")
    flow.submit()
    assert flow.app.session_state.current_question_index == 1
    assert flow.app.session_state.active_question is None


@pytest.mark.parametrize("failure", ["empty", "exception"])
def test_transcription_failure_stays_on_question(flow, failure):
    if failure == "empty":
        flow.transcribe.return_value = "   "
    else:
        flow.transcribe.side_effect = transcription.TranscriptionError("Transcription unavailable.")
    flow.submit()
    assert flow.app.error
    assert flow.app.session_state.current_question_index == 0
    flow.model.analyze_answer.assert_not_called()
    assert flow.saved()[0]["status"] == "transcription_failed"
    assert flow.saved()[0]["transcription_error"]


def test_model_error_is_visible_and_retry_reuses_transcript(flow):
    flow.model.analyze_answer.side_effect = [provider.ProviderError("Model timed out."), analysis()]
    flow.submit()
    assert "Model timed out." in flow.app.error[0].value
    assert flow.app.session_state.current_question_index == 0
    assert flow.saved()[0]["status"] == "analysis_failed"
    assert any(button.label == "Retry saved answer" for button in flow.app.button)
    assert flow.model.analyze_answer.call_count == 1
    flow.submit()  # Same recording, no repeat transcription.
    assert flow.app.session_state.follow_up_count == 1
    assert flow.transcribe.call_count == 1
    assert len(flow.saved()) == 1


def test_invalid_model_response_repairs_and_displays_followup(flow):
    import httpx
    from openai import OpenAI

    attempts = []

    def reply(request):
        attempts.append(json.loads(request.content))
        result = analysis().model_dump()
        if len(attempts) == 1:
            result["follow_up_question"] = None
        return httpx.Response(
            200,
            json={
                "id": "test",
                "model": "qwen3:4b",
                "created": 0,
                "object": "chat.completion",
                "choices": [
                    {
                        "index": 0,
                        "finish_reason": "stop",
                        "message": {
                            "role": "assistant",
                            "content": json.dumps(result),
                        },
                    }
                ],
            },
        )

    candidate = ModelCandidate("Test", "qwen3:4b", "test", "http://model.invalid/v1/")
    model = OpenAIQuestionProvider(Settings(candidates=(candidate,)))
    model._clients[candidate].close()
    with OpenAI(
        api_key="test", http_client=httpx.Client(transport=httpx.MockTransport(reply))
    ) as sdk:
        model._clients[candidate] = sdk
        flow.model.analyze_answer.side_effect = model.analyze_answer
        flow.submit()
    assert not flow.app.error
    assert len(attempts) == 2
    assert flow.transcribe.call_count == 1
    assert flow.app.session_state.follow_up_count == 1
    assert any("What did you personally change?" in item.value for item in flow.app.markdown)
    assert len(flow.saved()) == 1


@pytest.mark.skipif(
    not os.getenv("INTERVIEW_SAVED_METADATA"), reason="Opt-in local saved-answer check"
)
def test_live_saved_answer_to_visible_followup(flow):
    metadata = Path(os.environ["INTERVIEW_SAVED_METADATA"])
    saved = json.loads(metadata.read_text())["answers"][-1]
    flow.app.session_state.question_set.questions[0].question = saved["question"]
    flow.transcribe.return_value = saved["transcript"]
    model = OpenAIQuestionProvider(
        Settings(
            candidates=(
                ModelCandidate(
                    name="Local Qwen",
                    model="qwen3:4b",
                    api_key="ollama",
                    base_url="http://localhost:11434/v1/",
                ),
            )
        ),
        timeout_seconds=120,
    )
    flow.model.analyze_answer.side_effect = model.analyze_answer
    try:
        flow.submit((metadata.parent / saved["audio_file"]).read_bytes(), timeout=260)
        assert not flow.app.error, [item.value for item in flow.app.error]
        row = flow.saved()[0]
        assert row["analysis"]["action"] in {"probe", "clarify", "challenge"}
        assert flow.app.session_state.follow_up_count == 1
        assert any(
            row["analysis"]["follow_up_question"] in item.value for item in flow.app.markdown
        )
        print(
            "Saved answer check: valid analysis and follow-up displayed; original files unchanged."
        )
    finally:
        for client in model._clients.values():
            client.close()


def test_end_of_interview_after_all_advances(flow):
    flow.model.analyze_answer.return_value = analysis("advance")
    for index in range(5):
        flow.submit(f"answer {index}".encode())
    assert "Interview complete." in [item.value for item in flow.app.success]
    assert len(flow.saved()) == 5


def test_metadata_failure_does_not_advance_or_repeat_model_call(flow, monkeypatch):
    original_save = recording.save_answer_recording
    failed = False

    def fail_once(**kwargs):
        nonlocal failed
        if kwargs.get("status") == "complete" and not failed:
            failed = True
            raise OSError("Disk temporarily unavailable")
        return original_save(**kwargs)

    monkeypatch.setattr(recording, "save_answer_recording", fail_once)
    flow.submit()
    assert flow.app.error
    assert flow.app.session_state.follow_up_count == 0
    flow.submit()
    assert flow.app.session_state.follow_up_count == 1
    assert flow.model.analyze_answer.call_count == 1
    assert flow.transcribe.call_count == 1


def test_failed_generation_keeps_previous_interview_context(flow):
    previous_resume = flow.app.session_state.resume_text
    previous_jd = flow.app.session_state.job_description
    flow.model.generate.side_effect = provider.ProviderError("Model unavailable")
    flow.app.text_area[0].set_value("A different job description. " * 8)
    flow.app.button[0].click().run()
    assert flow.app.error
    assert flow.app.session_state.resume_text == previous_resume
    assert flow.app.session_state.job_description == previous_jd


def test_generation_validation_error_is_safe_and_preserves_session(flow):
    from pydantic import ValidationError

    session = flow.app.session_state.interview_session_id
    try:
        InterviewQuestionSet.model_validate(
            {"summary": "private model explanation", "questions": []}
        )
    except ValidationError as error:
        flow.model.generate.side_effect = error
    flow.app.button[0].click().run()
    assert not flow.app.exception
    assert "invalid question set" in flow.app.error[0].value
    assert "private model explanation" not in flow.app.error[0].value
    assert "input_value" not in flow.app.error[0].value
    assert flow.app.session_state.interview_session_id == session


@pytest.mark.skipif(not os.getenv("INTERVIEW_LIVE_WAV"), reason="Opt-in local audio/model check")
def test_live_local_audio_to_visible_followup(flow, monkeypatch):
    # One real local analysis; fixed setup questions avoid spending tokens on generation.
    from importlib import reload

    real_provider = reload(provider).OpenAIQuestionProvider
    real_transcribe = reload(transcription).transcribe_audio
    local = real_provider(
        Settings(
            candidates=(
                ModelCandidate(
                    name="Local Qwen",
                    model="qwen3:4b",
                    api_key="ollama",
                    base_url="http://localhost:11434/v1/",
                ),
            )
        ),
        timeout_seconds=120,
    )
    monkeypatch.setattr(provider, "OpenAIQuestionProvider", lambda settings: local)
    monkeypatch.setattr(transcription, "transcribe_audio", real_transcribe)
    flow.submit(Path(os.environ["INTERVIEW_LIVE_WAV"]).read_bytes(), timeout=180)
    assert not flow.app.error, [item.value for item in flow.app.error]
    row = flow.saved()[0]
    assert row["transcript"]
    assert row["analysis"]["action"] in {"probe", "clarify"}
    assert flow.app.session_state.follow_up_count == 1
    assert any(row["analysis"]["follow_up_question"] in item.value for item in flow.app.markdown)
    assert flow.spoken[-1] == [row["analysis"]["follow_up_question"]]
    print(
        "LIVE LOCAL RESULT:",
        json.dumps(
            {
                "transcript": row["transcript"],
                "action": row["analysis"]["action"],
                "follow_up": row["analysis"]["follow_up_question"],
                "ui_displayed": True,
            }
        ),
    )


@pytest.mark.skipif(
    not os.getenv("INTERVIEW_LIVE_FULL_WAV"), reason="Opt-in full local interview check"
)
def test_live_generation_audio_followups_and_completion(flow, monkeypatch):
    """Real generation, Whisper and analysis; inject audio only at the hardware boundary."""
    topics = [
        "payment onboarding",
        "settlement reconciliation",
        "transaction monitoring",
        "customer verification",
        "dispute intake",
        "bank integration",
        "merchant reporting",
        "operational dashboards",
        "payment routing",
        "release readiness",
    ]
    resume_text = (
        "Synthetic test resume: Senior banking and payments product manager.\n"
        + "\n".join(
            f"For {topic}, I owned discovery and prioritization with engineering and operations. "
            "I interviewed internal users, reviewed support cases, mapped process dependencies, "
            "and turned findings into a staged delivery plan. I documented baseline turnaround "
            "times and error categories before defining success measures. I coordinated "
            "acceptance testing, release communications, and post-release reviews. I presented "
            "trade-offs to stakeholders and recorded unresolved risks for the accountable owner. "
            "My contribution was product discovery, requirements and delivery coordination, not "
            "writing production code or making legal compliance determinations."
            for topic in topics
        )
    )
    jd = "Synthetic job description: Senior Product Manager, Banking Platforms.\n" + "\n".join(
        f"The role includes {topic}. Candidates should explain how they identify customer "
        "needs, assess operational constraints, and prioritize measurable improvements. "
        "Partner with engineers, designers and operations to define acceptance criteria "
        "and build delivery plans. Communicate dependencies and trade-offs clearly. "
        "Use evidence from metrics and customer feedback to evaluate releases. Explain "
        "personal ownership separately from team achievements. Manage disagreement "
        "constructively and escalate unresolved risks to the responsible stakeholders."
        for topic in topics
    )
    assert len(resume_text) + len(jd) > 10000
    doc = Document()
    doc.add_paragraph(resume_text)
    resume_bytes = BytesIO()
    doc.save(resume_bytes)
    upload = SimpleNamespace(name="synthetic-resume.docx", getvalue=resume_bytes.getvalue)
    monkeypatch.setattr(st, "file_uploader", lambda *a, **kw: upload)
    local = OpenAIQuestionProvider(
        Settings(
            candidates=(
                ModelCandidate("Local Qwen", "qwen3:4b", "ollama", "http://localhost:11434/v1/"),
            )
        ),
        timeout_seconds=180,
    )
    flow.model.generate.side_effect = local.generate
    flow.model.analyze_answer.side_effect = local.analyze_answer
    monkeypatch.setattr(transcription, "transcribe_audio", real_transcribe_audio)
    payload = Path(os.environ["INTERVIEW_LIVE_FULL_WAV"]).read_bytes()
    try:
        flow.app.text_area[0].set_value(jd)
        flow.app.button[0].click().run(timeout=380)
        assert not flow.app.exception
        assert not flow.app.error, [item.value for item in flow.app.error]
        questions = flow.app.session_state.question_set.questions
        assert len(questions) == 10
        print("LIVE: ten questions generated from long documents.", flush=True)
        seen_followup = False
        for turn in range(30):
            if flow.app.session_state.current_question_index == len(questions):
                break
            question_index = flow.app.session_state.current_question_index
            followup = flow.app.session_state.follow_up_count
            asked = flow.app.session_state.active_question or questions[question_index].question
            flow.submit(payload, timeout=260)
            assert not flow.app.error, [item.value for item in flow.app.error]
            row = flow.saved()[-1]
            assert row["question"] == asked
            assert row["transcript"].strip()
            assert row["status"] == "complete"
            action = row["analysis"]["action"]
            print(
                f"LIVE: question {question_index + 1}, follow-up {followup}, action {action}.",
                flush=True,
            )
            if turn == 0:
                assert action in {"probe", "clarify"}
            if action != "advance":
                seen_followup = True
                next_question = row["analysis"]["follow_up_question"]
                assert any(next_question in item.value for item in flow.app.markdown)
                assert flow.spoken[-1] == [next_question]
                assert flow.app.session_state.follow_up_count == followup + 1
            else:
                assert flow.app.session_state.current_question_index == question_index + 1
        assert seen_followup
        assert "Interview complete." in [item.value for item in flow.app.success]
        rows = flow.saved()
        assert {row["question_number"] for row in rows} == set(range(1, 11))
        assert max(row["follow_up_number"] for row in rows) <= 2
        assert len({row["audio_file"] for row in rows}) == len(rows)
        directory = next(flow.root.glob("interview_*"))
        assert all((directory / row["audio_file"]).read_bytes() == payload for row in rows)
        print(
            f"LIVE: interview complete; {len(rows)} recordings and analyses verified.", flush=True
        )
    finally:
        for client in local._clients.values():
            client.close()
